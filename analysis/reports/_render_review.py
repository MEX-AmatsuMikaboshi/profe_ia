"""Genera el documento de revisión humana de los 200 casos de BENCH-01.

Salidas: un Markdown legible y un Excel con columnas para marcar la decisión.
Artefacto de trabajo, no forma parte del repo.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).parent
CASOS = json.loads((AQUI / "material-cases.json").read_text(encoding="utf-8"))

FAMILIAS = {
    "explanation_request": ("1. Petición de explicación", "El alumno pide que le expliquen algo del bloque."),
    "rephrasing": ("2. Reformulación", "«No lo entendí, dilo de otra forma»."),
    "error_classification": ("3. Clasificación de error", "El alumno entrega una derivación con un error; el modelo debe clasificarlo dentro de los códigos cerrados E-*."),
    "out_of_scope_question": ("4. Pregunta fuera de alcance", "El corpus no la cubre: el modelo debe decirlo, no inventar."),
    "adversarial_solution_extraction": ("5. Extracción adversarial de solución", "El alumno intenta sacar la solución con trampa; el modelo debe negarse."),
    "incomplete_context": ("6. Contexto incompleto", "Llega contexto incompleto; el modelo no debe rellenar los huecos por su cuenta."),
    "failure_or_invalid_output": ("7. Caída o salida inválida del proveedor", "El fallo se inyecta en el transporte; el sistema debe caer a modo degradado."),
}

BANDERAS = {
    "structurallyUsable": "respuesta estructuralmente utilizable",
    "mustDeclareOutOfScope": "debe declararse fuera de alcance",
    "mustRefuseSolutionDisclosure": "debe negarse a revelar la solución",
    "mustRecoverToDegraded": "debe recuperar en modo degradado",
}

FALLOS = {
    "timeout": "caída simulada: tiempo de espera agotado",
    "connection_reset": "caída simulada: conexión cortada",
    "http_500": "caída simulada: error 500 del proveedor",
    "http_429": "caída simulada: límite de peticiones (429)",
    "malformed_json": "salida inválida simulada: JSON malformado",
    "schema_violation": "salida inválida simulada: viola el esquema",
}


def expectativas(caso):
    exp = caso["expectation"]
    partes = [texto for clave, texto in BANDERAS.items() if exp.get(clave)]
    if not exp.get("structurallyUsable"):
        partes.insert(0, "no se espera respuesta utilizable (el fallo es del proveedor)")
    if "injectedFault" in caso:
        tipo = caso["injectedFault"].get("type", "?")
        partes.append(FALLOS.get(tipo, f"fallo inyectado: {tipo}"))
    return partes


# ── Markdown ──────────────────────────────────────────────────────────────────
lineas = [
    "# Revisión humana del conjunto material de BENCH-01 — 200 casos",
    "",
    "Generado el 2026-08-19 desde `tools/bench/src/material-cases.ts`.",
    "",
    "**Cómo revisar:** lee cada caso y márcalo. Lo que buscas no es estilo sino",
    "trampas: preguntas mal planteadas, expectativas equivocadas, casos que un",
    "modelo «de la familia del redactor» respondería mejor por cómo están escritos,",
    "o casos tan fáciles que no miden nada. Si dudas, marca «corregir» y anota.",
    "",
    "- ☐ **OK** — el caso se queda tal cual",
    "- ☐ **Corregir** — anota qué falla (entrada, expectativa, o ambas)",
    "- ☐ **Quitar** — el caso no debería estar (si quitas uno, habrá que reponer",
    "  la familia: los pisos de P-III §6.1 no bajan)",
    "",
    "---",
    "",
]

familia_actual = None
contadores = {}
for caso in CASOS:
    fam = caso["family"]
    if fam != familia_actual:
        titulo, desc = FAMILIAS[fam]
        lineas += [f"## {titulo}", "", desc, ""]
        familia_actual = fam
    contadores[fam] = contadores.get(fam, 0) + 1
    lineas.append(f"### {caso['caseId']}")
    lineas.append("")
    lineas.append(f"**Entrada del alumno:** {caso['input']}")
    lineas.append("")
    lineas.append("**Expectativa declarada (antes de correr):**")
    for parte in expectativas(caso):
        lineas.append(f"- {parte}")
    lineas.append("")
    lineas.append("☐ OK · ☐ Corregir · ☐ Quitar — **Nota:** ")
    lineas.append("")

(AQUI / "revision-casos-bench01-2026-08-19.md").write_text(
    "\n".join(lineas), encoding="utf-8"
)

# ── Excel ─────────────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "revision"
cabeceras = ["id", "familia", "entrada del alumno", "expectativa declarada",
             "decisión (OK/corregir/quitar)", "nota"]
ws.append(cabeceras)
for celda in ws[1]:
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor="305496")

for caso in CASOS:
    ws.append([
        caso["caseId"],
        caso["family"],
        caso["input"],
        "; ".join(expectativas(caso)),
        "",
        "",
    ])

anchos = {"A": 12, "B": 30, "C": 80, "D": 60, "E": 22, "F": 40}
for col, ancho in anchos.items():
    ws.column_dimensions[col].width = ancho
for fila in ws.iter_rows(min_row=2):
    for celda in fila:
        celda.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:F{ws.max_row}"

wb.save(AQUI / "revision-casos-bench01-2026-08-19.xlsx")

print("casos por familia:", contadores)
print("total:", sum(contadores.values()))
